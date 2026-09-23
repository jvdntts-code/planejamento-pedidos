import hashlib
import io
import re
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


NAVY = "17365D"
WHITE = "FFFFFF"


def br_to_float(value):
    if value is None:
        return 0.0
    text = str(value).strip().replace(".", "").replace(",", ".")
    try:
        return float(text)
    except Exception:
        return 0.0


def money_br(value):
    return (
        f"R$ {float(value):,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


def extract_pdf_text(uploaded):
    raw = uploaded if isinstance(uploaded, (bytes, bytearray)) else uploaded.getvalue()
    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    full_text = "\n".join(pages).strip()
    if not full_text:
        raise ValueError(
            "Não consegui extrair texto deste PDF. "
            "Este primeiro modelo funciona com PDFs gerados pelo sistema, que possuem texto selecionável."
        )
    return full_text, len(reader.pages)


def _smart_join_item_lines(lines):
    """Reconstrói itens que quebraram em duas ou três linhas durante a leitura do PDF."""
    rebuilt = []
    i = 0

    while i < len(lines):
        line = lines[i].strip()
        if not re.match(r"^PC\d+", line, flags=re.IGNORECASE):
            i += 1
            continue

        buffer = line
        j = i + 1

        while (
            not re.search(r"\d{6}$", buffer)
            and j < len(lines)
            and not re.match(r"^PC\d+", lines[j].strip(), flags=re.IGNORECASE)
            and "TOTAL BRUTO" not in lines[j].upper()
        ):
            next_line = lines[j].strip()

            # Corrige palavras quebradas pelo PDF, por exemplo:
            # "TERMINA" + "IS" -> "TERMINAIS"; "FURO" + "S" -> "FUROS".
            if (
                re.fullmatch(r"[A-Za-zÀ-ÿ]{1,3}", next_line)
                and re.search(r"[A-Za-zÀ-ÿ]$", buffer)
            ):
                buffer += next_line
            else:
                buffer += " " + next_line
            j += 1

        rebuilt.append(buffer)
        i = j

    return rebuilt


def parse_items(text):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    item_lines = _smart_join_item_lines(lines)

    # Formato observado no PDF do sistema:
    # PC25 0,0044099001 36,38 0,00 DESCRICAO000139
    #   PC = unidade
    #   25 = quantidade
    #   0,00 = IPI
    #   44099 = número do fabricante
    #   001 = número do item
    #   36,38 = valor unitário
    #   0,00 = desconto
    #   000139 = código interno
    pattern = re.compile(
        r"^PC(?P<qtd>\d+(?:[.,]\d+)?)\s+"
        r"(?P<ipi>\d+,\d{2})"
        r"(?P<fabricante_item>[A-Z0-9]+?)"
        r"(?P<item>\d{3})\s+"
        r"(?P<valor_unitario>\d+,\d{2})\s+"
        r"(?P<desconto>\d+,\d{2})\s+"
        r"(?P<descricao>.+?)"
        r"(?P<codigo>\d{6})$",
        flags=re.IGNORECASE,
    )

    rows = []
    not_parsed = []

    for raw_line in item_lines:
        match = pattern.match(raw_line)
        if not match:
            not_parsed.append(raw_line)
            continue

        data = match.groupdict()
        fabricante_item = data["fabricante_item"].strip()
        qtd = br_to_float(data["qtd"])
        valor_unitario = br_to_float(data["valor_unitario"])
        desconto = br_to_float(data["desconto"])
        ipi = br_to_float(data["ipi"])

        rows.append(
            {
                "Item": int(data["item"]),
                "Codigo": data["codigo"],
                "Nr Fabricante": fabricante_item,
                "Descricao": re.sub(r"\s+", " ", data["descricao"]).strip(),
                "Qtd Pedida": qtd,
                "D/U": "PC",
                "Vl Unitario": valor_unitario,
                "% Desc": desconto,
                "% IPI": ipi,
                "Valor Item": qtd * valor_unitario,
                "Qtd Recebida": 0.0,
                "Qtd Pendente": qtd,
                "Status": "EM ABERTO",
            }
        )

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values("Item").reset_index(drop=True)

    return df, not_parsed


def parse_order_header(text):
    supplier_match = re.search(
        r"\b(F\d{6})\s+([A-ZÀ-ÿ0-9 .&/-]+?LTDA)\b",
        text,
        flags=re.IGNORECASE,
    )
    order_match = re.search(r"S/PEDIDO\.*:\s*(\d+)", text, flags=re.IGNORECASE)

    date_match = re.search(
        r"DATA PEDIDO\.*:\s*(?:\S+\s*)?(\d{2}/\d{2}/\d{4})",
        text,
        flags=re.IGNORECASE,
    )
    if not date_match:
        date_match = re.search(
            r"EMISS[ÃA]O\s*:\s*(\d{2}/\d{2}/\d{4})",
            text,
            flags=re.IGNORECASE,
        )

    total_match = re.search(
        r"VALOR PEDIDO\.*:\s*(?:% DESC\s*)?([\d.]+,\d{2})",
        text,
        flags=re.IGNORECASE,
    )
    if not total_match:
        total_match = re.search(
            r"TOTAL LIQUIDO DO PEDIDO-->\s*(?:.*?\n)*?([\d.]+,\d{2})",
            text,
            flags=re.IGNORECASE | re.DOTALL,
        )

    buyer_match = re.search(
        r"COMPRADOR\.*:\s*([^\n]+)",
        text,
        flags=re.IGNORECASE,
    )

    supplier_code = supplier_match.group(1).upper() if supplier_match else ""
    supplier_name = supplier_match.group(2).strip().upper() if supplier_match else ""
    order_number = order_match.group(1) if order_match else ""
    order_date = date_match.group(1) if date_match else ""
    total_value = br_to_float(total_match.group(1)) if total_match else 0.0
    buyer = buyer_match.group(1).strip() if buyer_match else ""

    return {
        "Fornecedor Codigo": supplier_code,
        "Fornecedor": supplier_name,
        "Pedido": order_number,
        "Data Pedido": order_date,
        "Valor Pedido": total_value,
        "Comprador": buyer,
    }


def calculate_next_cycle(order_date, months=2):
    try:
        dt = datetime.strptime(order_date, "%d/%m/%Y")
    except Exception:
        return None
    return (pd.Timestamp(dt) + pd.DateOffset(months=months)).date()


def parse_order_pdf(uploaded):
    text, pages = extract_pdf_text(uploaded)
    header = parse_order_header(text)
    items, not_parsed = parse_items(text)

    if items.empty:
        raise ValueError(
            "O PDF foi lido, mas nenhum item foi reconhecido. "
            "O layout pode ser diferente do modelo de Pedido Fornecedor usado na configuração inicial."
        )

    header["Paginas"] = pages
    header["Itens"] = len(items)
    header["Quantidade Total"] = float(items["Qtd Pedida"].sum())
    header["Valor Calculado Itens"] = float(items["Valor Item"].sum())
    header["Proxima Compra"] = None
    header["Ciclo Meses"] = None

    return header, items, not_parsed, text


def update_receipt_columns(df):
    x = df.copy()
    x["Qtd Pedida"] = pd.to_numeric(x["Qtd Pedida"], errors="coerce").fillna(0)
    x["Qtd Recebida"] = pd.to_numeric(x["Qtd Recebida"], errors="coerce").fillna(0)
    x["Qtd Recebida"] = x[["Qtd Recebida", "Qtd Pedida"]].min(axis=1).clip(lower=0)
    x["Qtd Pendente"] = (x["Qtd Pedida"] - x["Qtd Recebida"]).clip(lower=0)

    x["Status"] = "EM ABERTO"
    x.loc[x["Qtd Recebida"] > 0, "Status"] = "PARCIAL"
    x.loc[x["Qtd Pendente"] <= 0, "Status"] = "RECEBIDO"
    return x


def export_order_xlsx(header, items):
    summary = pd.DataFrame(
        {
            "Campo": [
                "Fornecedor",
                "Código Fornecedor",
                "Pedido",
                "Data Pedido",
                "Ciclo entre compras (meses)",
                "Próxima Compra",
                "Valor Pedido",
                "Quantidade de Itens",
                "Quantidade Total",
                "Valor Calculado dos Itens",
            ],
            "Valor": [
                header.get("Fornecedor", ""),
                header.get("Fornecedor Codigo", ""),
                header.get("Pedido", ""),
                header.get("Data Pedido", ""),
                header.get("Ciclo Meses", ""),
                (
                    header.get("Proxima Compra").strftime("%d/%m/%Y")
                    if header.get("Proxima Compra") is not None
                    else ""
                ),
                header.get("Valor Pedido", 0),
                len(items),
                float(items["Qtd Pedida"].sum()) if not items.empty else 0,
                float(items["Valor Item"].sum()) if not items.empty else 0,
            ],
        }
    )

    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="PEDIDO", index=False)
        items.to_excel(writer, sheet_name="ITENS", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions

            for cell in ws[1]:
                cell.font = Font(name="Times New Roman", bold=True, color=WHITE)
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Times New Roman", size=11)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                max_len = max([len(str(cell.value or "")) for cell in col[:400]] + [10])
                ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 46)

    out.seek(0)
    return out.getvalue()


def _order_attachment_id(file_bytes):
    return hashlib.sha256(file_bytes).hexdigest()[:16]


def _ensure_order_library():
    if "gestao_pedidos_anexos" not in st.session_state:
        st.session_state["gestao_pedidos_anexos"] = {}
    if "gestao_pedido_selecionado" not in st.session_state:
        st.session_state["gestao_pedido_selecionado"] = None
    return st.session_state["gestao_pedidos_anexos"]


def _add_order_attachments(uploaded_files):
    library = _ensure_order_library()
    added = 0
    errors = []

    for uploaded in uploaded_files or []:
        raw = uploaded.getvalue()
        attachment_id = _order_attachment_id(raw)

        if attachment_id in library:
            continue

        try:
            header, items, not_parsed, extracted_text = parse_order_pdf(raw)
            library[attachment_id] = {
                "id": attachment_id,
                "filename": uploaded.name,
                "bytes": raw,
                "header": header,
                "items": items,
                "not_parsed": not_parsed,
                "extracted_text": extracted_text,
            }
            added += 1
        except Exception as exc:
            errors.append(f"{uploaded.name}: {exc}")

    return added, errors


def _render_order_detail(order_data, ciclo_meses):
    header = dict(order_data["header"])
    items = order_data["items"].copy()
    not_parsed = order_data["not_parsed"]
    extracted_text = order_data["extracted_text"]

    header["Ciclo Meses"] = int(ciclo_meses)
    header["Proxima Compra"] = calculate_next_cycle(
        header.get("Data Pedido", ""),
        int(ciclo_meses),
    )

    top_left, top_right = st.columns([1, 5])
    with top_left:
        if st.button(
            "← Voltar",
            use_container_width=True,
            key=f"voltar_pedidos_{order_data['id']}",
        ):
            st.session_state["gestao_pedido_selecionado"] = None
            st.rerun()

    with top_right:
        st.caption(f"📎 Anexo: {order_data['filename']}")

    st.success(
        f"Pedido reconhecido: {header.get('Pedido') or 'sem número'} | "
        f"{header.get('Itens', 0)} item(ns) extraído(s)."
    )

    if not_parsed:
        st.warning(
            f"{len(not_parsed)} linha(s) com aparência de item não foram interpretadas. "
            "Confira a tabela antes de continuar."
        )

    st.markdown("### Visão geral")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Pedido", header.get("Pedido") or "—")
    c2.metric("Fornecedor", header.get("Fornecedor") or "—")
    c3.metric("Valor do pedido", money_br(header.get("Valor Pedido", 0)))
    c4.metric("Itens", int(header.get("Itens", 0)))

    c5, c6, c7, c8 = st.columns(4)
    c5.metric("Data do pedido", header.get("Data Pedido") or "—")
    next_cycle = header.get("Proxima Compra")
    c6.metric(
        f"Próxima compra ({int(ciclo_meses)} {'mês' if int(ciclo_meses) == 1 else 'meses'})",
        next_cycle.strftime("%d/%m/%Y") if next_cycle is not None else "—",
    )
    c7.metric(
        "Quantidade total",
        f"{header.get('Quantidade Total', 0):,.0f}".replace(",", "."),
    )
    c8.metric(
        "Valor calculado pelos itens",
        money_br(header.get("Valor Calculado Itens", 0)),
    )

    pdf_total = float(header.get("Valor Pedido", 0) or 0)
    item_total = float(header.get("Valor Calculado Itens", 0) or 0)
    if pdf_total and abs(pdf_total - item_total) <= 0.02:
        st.success("✅ O valor somado dos itens confere com o valor total do PDF.")
    elif pdf_total:
        st.warning(
            "⚠️ O valor somado dos itens não confere com o total informado no PDF. "
            f"Diferença: {money_br(item_total - pdf_total)}."
        )

    st.markdown("### Itens do pedido")
    st.caption(
        "Você pode preencher a coluna **Qtd Recebida**. "
        "O saldo pendente e o status são recalculados logo abaixo."
    )

    editor_key = f"order_items_{order_data['id']}"
    edited = st.data_editor(
        items,
        use_container_width=True,
        hide_index=True,
        key=editor_key,
        disabled=[
            "Item",
            "Codigo",
            "Nr Fabricante",
            "Descricao",
            "Qtd Pedida",
            "D/U",
            "Vl Unitario",
            "% Desc",
            "% IPI",
            "Valor Item",
            "Qtd Pendente",
            "Status",
        ],
        column_config={
            "Qtd Recebida": st.column_config.NumberColumn(
                "Qtd Recebida",
                min_value=0.0,
                step=1.0,
            ),
            "Vl Unitario": st.column_config.NumberColumn(
                "Vl Unitario",
                format="R$ %.2f",
            ),
            "Valor Item": st.column_config.NumberColumn(
                "Valor Item",
                format="R$ %.2f",
            ),
        },
    )

    managed = update_receipt_columns(edited)
    st.session_state["gestao_pedidos_anexos"][order_data["id"]]["items"] = managed.copy()

    received_qty = float(managed["Qtd Recebida"].sum())
    pending_qty = float(managed["Qtd Pendente"].sum())
    received_items = int((managed["Status"] == "RECEBIDO").sum())
    partial_items = int((managed["Status"] == "PARCIAL").sum())

    r1, r2, r3, r4 = st.columns(4)
    r1.metric("Qtd recebida", f"{received_qty:,.0f}".replace(",", "."))
    r2.metric("Qtd pendente", f"{pending_qty:,.0f}".replace(",", "."))
    r3.metric("Itens recebidos", received_items)
    r4.metric("Itens parciais", partial_items)

    with st.expander("Ver acompanhamento atualizado"):
        st.dataframe(
            managed[
                [
                    "Item",
                    "Codigo",
                    "Nr Fabricante",
                    "Descricao",
                    "Qtd Pedida",
                    "Qtd Recebida",
                    "Qtd Pendente",
                    "Status",
                    "Vl Unitario",
                    "Valor Item",
                ]
            ],
            use_container_width=True,
            hide_index=True,
        )

    st.download_button(
        "📥 Exportar pedido para Excel",
        data=export_order_xlsx(header, managed),
        file_name=f"NEXO_pedido_{header.get('Pedido') or 'importado'}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
        key=f"export_order_xlsx_{order_data['id']}",
    )

    with st.expander("📄 Ver PDF anexado"):
        st.download_button(
            "📎 Baixar PDF original",
            data=order_data["bytes"],
            file_name=order_data["filename"],
            mime="application/pdf",
            key=f"download_pdf_{order_data['id']}",
        )

    with st.expander("🔧 Diagnóstico da leitura do PDF"):
        st.caption(
            "Use esta área apenas se algum dado não for reconhecido corretamente em outro modelo de pedido."
        )
        st.text_area(
            "Texto extraído",
            extracted_text,
            height=220,
            disabled=True,
            key=f"order_extracted_text_{order_data['id']}",
        )


def render_gestao_pedidos():
    st.title("🧾 Gestão de Pedidos")
    st.caption(
        "Centralize seus pedidos em PDF e abra cada pedido para acompanhar itens, recebimentos e próxima compra."
    )

    library = _ensure_order_library()

    with st.sidebar:
        st.markdown("---")
        st.markdown("### Gestão de Pedidos")
        ciclo_meses = st.number_input(
            "Intervalo entre compras (meses)",
            min_value=1,
            max_value=24,
            value=2,
            step=1,
            key="gestao_ciclo_meses",
            help=(
                "Define quantos meses após a data do pedido o NEXO deve indicar "
                "como próxima compra."
            ),
        )
        st.caption(
            f"Próxima compra será projetada {int(ciclo_meses)} "
            f"{'mês' if int(ciclo_meses) == 1 else 'meses'} após cada pedido."
        )

    selected_id = st.session_state.get("gestao_pedido_selecionado")
    if selected_id and selected_id in library:
        _render_order_detail(library[selected_id], ciclo_meses)
        return
    elif selected_id:
        st.session_state["gestao_pedido_selecionado"] = None

    st.markdown("### Pedidos anexados")

    with st.expander("➕ Anexar pedido(s) em PDF", expanded=not bool(library)):
        st.caption(
            "Você pode selecionar vários PDFs de uma vez. Cada arquivo será transformado em um pedido na lista abaixo."
        )
        uploaded_files = st.file_uploader(
            "Selecionar Pedido(s) Fornecedor em PDF",
            type=["pdf"],
            accept_multiple_files=True,
            key="gestao_pedidos_pdf_multiplos",
            help="Use os PDFs originais gerados pelo sistema.",
        )

        if uploaded_files:
            added, errors = _add_order_attachments(uploaded_files)
            if added:
                st.success(f"{added} pedido(s) adicionado(s) à tela.")
            for error in errors:
                st.error(error)

    if not library:
        st.info("Ainda não há pedidos anexados. Use a opção acima para adicionar o primeiro PDF.")
        return

    search = st.text_input(
        "Buscar pedido ou fornecedor",
        key="gestao_busca_pedidos",
        placeholder="Ex.: 7432 ou VALCLEI",
    ).strip().lower()

    records = []
    for attachment_id, order_data in library.items():
        header = order_data["header"]
        records.append(
            {
                "id": attachment_id,
                "pedido": str(header.get("Pedido") or ""),
                "fornecedor": str(header.get("Fornecedor") or ""),
                "data": str(header.get("Data Pedido") or ""),
                "valor": float(header.get("Valor Pedido", 0) or 0),
                "itens": int(header.get("Itens", 0) or 0),
                "filename": order_data["filename"],
            }
        )

    def date_sort_value(record):
        try:
            return datetime.strptime(record["data"], "%d/%m/%Y")
        except Exception:
            return datetime.min

    records.sort(key=date_sort_value, reverse=True)

    if search:
        records = [
            rec for rec in records
            if search in rec["pedido"].lower()
            or search in rec["fornecedor"].lower()
            or search in rec["filename"].lower()
        ]

    total_value = sum(rec["valor"] for rec in records)
    k1, k2, k3 = st.columns(3)
    k1.metric("Pedidos anexados", len(records))
    k2.metric("Valor listado", money_br(total_value))
    k3.metric("Arquivos PDF", len(records))

    if not records:
        st.warning("Nenhum pedido encontrado para essa busca.")
        return

    meses_pt = {
        1: "JANEIRO", 2: "FEVEREIRO", 3: "MARÇO", 4: "ABRIL",
        5: "MAIO", 6: "JUNHO", 7: "JULHO", 8: "AGOSTO",
        9: "SETEMBRO", 10: "OUTUBRO", 11: "NOVEMBRO", 12: "DEZEMBRO",
    }

    grupos = {}
    for rec in records:
        try:
            dt = datetime.strptime(rec["data"], "%d/%m/%Y")
            chave = (dt.year, dt.month)
            titulo_mes = f"{meses_pt[dt.month]} {dt.year}"
        except Exception:
            chave = (0, 0)
            titulo_mes = "SEM DATA"

        grupos.setdefault(
            chave,
            {
                "titulo": titulo_mes,
                "pedidos": [],
            },
        )["pedidos"].append(rec)

    grupos_ordenados = sorted(
        grupos.items(),
        key=lambda item: item[0],
        reverse=True,
    )

    for indice_grupo, (_, grupo) in enumerate(grupos_ordenados):
        pedidos_mes = grupo["pedidos"]
        valor_mes = sum(rec["valor"] for rec in pedidos_mes)
        qtd_mes = len(pedidos_mes)

        titulo_expander = (
            f"📅 {grupo['titulo']}  •  "
            f"{qtd_mes} {'pedido' if qtd_mes == 1 else 'pedidos'}  •  "
            f"{money_br(valor_mes)}"
        )

        with st.expander(
            titulo_expander,
            expanded=(indice_grupo == 0),
        ):
            for rec in pedidos_mes:
                with st.container(border=True):
                    col_main, col_meta, col_actions = st.columns([3.2, 2, 1.25])

                    with col_main:
                        st.markdown(f"### 📎 Pedido {rec['pedido'] or 'sem número'}")
                        st.markdown(
                            f"**{rec['fornecedor'] or 'Fornecedor não identificado'}**"
                        )
                        st.caption(rec["filename"])

                    with col_meta:
                        st.markdown(f"**Data:** {rec['data'] or '—'}")
                        st.markdown(f"**Valor:** {money_br(rec['valor'])}")
                        st.markdown(f"**Itens:** {rec['itens']}")

                    with col_actions:
                        if st.button(
                            "Abrir pedido",
                            use_container_width=True,
                            type="primary",
                            key=f"abrir_pedido_{rec['id']}",
                        ):
                            st.session_state["gestao_pedido_selecionado"] = rec["id"]
                            st.rerun()

                        if st.button(
                            "Remover",
                            use_container_width=True,
                            key=f"remover_pedido_{rec['id']}",
                        ):
                            library.pop(rec["id"], None)
                            st.rerun()

    st.caption(
        "Os PDFs anexados ficam disponíveis durante esta sessão do NEXO. "
        "Eles não são enviados ao repositório público do GitHub."
    )

