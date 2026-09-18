import io
import re
import unicodedata

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


NAVY = "17365D"
WHITE = "FFFFFF"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "F4CCCC"
LIGHT_ORANGE = "FCE4D6"

STATUS_ORDER = [
    "QUANTIDADE DIVERGENTE",
    "SÓ NO FORNECEDOR",
    "SÓ NO SISTEMA",
    "OK",
]

STATUS_FILL = {
    "QUANTIDADE DIVERGENTE": LIGHT_YELLOW,
    "SÓ NO FORNECEDOR": LIGHT_ORANGE,
    "SÓ NO SISTEMA": LIGHT_RED,
    "OK": LIGHT_GREEN,
}


def norm(value):
    value = str(value).strip().lower()
    value = "".join(
        c for c in unicodedata.normalize("NFKD", value)
        if not unicodedata.combining(c)
    )
    return re.sub(r"[^a-z0-9]+", "", value)


def normalize_reference(value):
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    text = str(value).strip().upper()
    text = re.sub(r"\.0$", "", text)
    text = re.sub(r"\s+", "", text)
    return text


def find_col(df, *names):
    mapping = {norm(c): c for c in df.columns}
    for name in names:
        key = norm(name)
        if key in mapping:
            return mapping[key]
    return None


def read_single_table(uploaded):
    name = uploaded.name.lower()
    raw = uploaded.getvalue()
    if name.endswith((".xlsx", ".xls")):
        return pd.read_excel(io.BytesIO(raw))
    if name.endswith(".csv"):
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                return pd.read_csv(io.BytesIO(raw), sep=None, engine="python", encoding=enc)
            except Exception:
                pass
    raise ValueError("Arquivo não reconhecido. Use XLSX, XLS ou CSV.")


def read_combined_workbook(uploaded):
    raw = uploaded.getvalue()
    book = pd.read_excel(io.BytesIO(raw), sheet_name=None)
    if not book:
        raise ValueError("O arquivo não possui abas legíveis.")

    supplier_sheet = None
    system_sheet = None
    for sheet_name in book:
        key = norm(sheet_name)
        if supplier_sheet is None and "fornecedor" in key:
            supplier_sheet = sheet_name
        if system_sheet is None and "sistema" in key:
            system_sheet = sheet_name

    if supplier_sheet is None or system_sheet is None:
        raise ValueError(
            "Não encontrei as duas abas esperadas. O arquivo deve ter uma aba do FORNECEDOR e outra do SISTEMA."
        )

    return book[system_sheet], book[supplier_sheet], system_sheet, supplier_sheet


def prepare_pending(df, source):
    if df is None or df.empty:
        raise ValueError(f"A base de {source.lower()} está vazia.")

    ref_col = find_col(
        df,
        "REFERENCIA",
        "REFERÊNCIA",
        "NumFabricante",
        "Numero Fabricante",
        "Referência Fabricante",
        "Codigo Fabricante",
    )

    if source == "SISTEMA":
        qty_col = find_col(
            df,
            "SOMA PENDENTE",
            "QTD PENDENTE",
            "PENDENCIA",
            "PENDÊNCIA",
            "QUANTIDADE PENDENTE",
            "QUANTIDADE",
            "QTD",
        )
    else:
        qty_col = find_col(
            df,
            "SOMA PENDENTE FORNECEDOR",
            "QTD PENDENTE FORNECEDOR",
            "PENDENCIA FORNECEDOR",
            "PENDÊNCIA FORNECEDOR",
            "QUANTIDADE PENDENTE",
            "QUANTIDADE",
            "QTD PENDENTE",
            "QTD",
        )

    if ref_col is None or qty_col is None:
        expected = (
            "REFERENCIA + SOMA PENDENTE"
            if source == "SISTEMA"
            else "REFERENCIA + SOMA PENDENTE FORNECEDOR"
        )
        raise ValueError(
            f"Não consegui identificar as colunas da base {source}. Esperado: {expected}."
        )

    out = pd.DataFrame()
    out["REFERENCIA"] = df[ref_col].map(normalize_reference)
    out["QUANTIDADE"] = pd.to_numeric(df[qty_col], errors="coerce").fillna(0)
    out = out[out["REFERENCIA"].ne("")].copy()
    out = out.groupby("REFERENCIA", as_index=False)["QUANTIDADE"].sum()
    return out


def compare_pending(system_df, supplier_df):
    system = prepare_pending(system_df, "SISTEMA").rename(
        columns={"QUANTIDADE": "QTD SISTEMA"}
    )
    supplier = prepare_pending(supplier_df, "FORNECEDOR").rename(
        columns={"QUANTIDADE": "QTD FORNECEDOR"}
    )

    system["TEM NO SISTEMA"] = True
    supplier["TEM NO FORNECEDOR"] = True

    result = system.merge(supplier, on="REFERENCIA", how="outer")
    result["TEM NO SISTEMA"] = result["TEM NO SISTEMA"].fillna(False).astype(bool)
    result["TEM NO FORNECEDOR"] = result["TEM NO FORNECEDOR"].fillna(False).astype(bool)
    result["QTD SISTEMA"] = pd.to_numeric(result["QTD SISTEMA"], errors="coerce").fillna(0)
    result["QTD FORNECEDOR"] = pd.to_numeric(result["QTD FORNECEDOR"], errors="coerce").fillna(0)
    result["DIFERENÇA FORNECEDOR - SISTEMA"] = result["QTD FORNECEDOR"] - result["QTD SISTEMA"]

    conditions = [
        (~result["TEM NO SISTEMA"]) & result["TEM NO FORNECEDOR"],
        result["TEM NO SISTEMA"] & (~result["TEM NO FORNECEDOR"]),
        result["TEM NO SISTEMA"]
        & result["TEM NO FORNECEDOR"]
        & (~np.isclose(result["QTD SISTEMA"], result["QTD FORNECEDOR"])),
    ]
    choices = [
        "SÓ NO FORNECEDOR",
        "SÓ NO SISTEMA",
        "QUANTIDADE DIVERGENTE",
    ]
    result["STATUS"] = np.select(conditions, choices, default="OK")

    def reading(row):
        if row["STATUS"] == "SÓ NO FORNECEDOR":
            return "Fornecedor possui pendência que não aparece no sistema"
        if row["STATUS"] == "SÓ NO SISTEMA":
            return "Sistema possui pendência que não aparece no fornecedor"
        if row["STATUS"] == "QUANTIDADE DIVERGENTE":
            diff = row["DIFERENÇA FORNECEDOR - SISTEMA"]
            if diff > 0:
                return f"Fornecedor possui {abs(diff):g} unidade(s) a mais"
            return f"Sistema possui {abs(diff):g} unidade(s) a mais"
        return "Quantidades conferem"

    result["LEITURA"] = result.apply(reading, axis=1)

    order_map = {name: idx for idx, name in enumerate(STATUS_ORDER)}
    result["_ORDEM"] = result["STATUS"].map(order_map).fillna(99)
    result = result.sort_values(["_ORDEM", "REFERENCIA"], kind="stable").drop(columns="_ORDEM")
    return result.reset_index(drop=True)


def pending_template_bytes():
    supplier = pd.DataFrame(columns=["REFERENCIA", "SOMA PENDENTE FORNECEDOR"])
    system = pd.DataFrame(columns=["REFERENCIA", "SOMA PENDENTE"])
    example_supplier = pd.DataFrame({
        "REFERENCIA": ["REF001", "REF002", "REF004"],
        "SOMA PENDENTE FORNECEDOR": [10, 5, 8],
    })
    example_system = pd.DataFrame({
        "REFERENCIA": ["REF001", "REF002", "REF003"],
        "SOMA PENDENTE": [10, 7, 12],
    })
    instructions = pd.DataFrame([
        {"Etapa": 1, "Orientação": "O arquivo pode conter as duas abas: PENDENCIA FORNECEDOR SOMA e PENDENCIA SISTEMA SOMA."},
        {"Etapa": 2, "Orientação": "Use REFERENCIA como chave de comparação entre o sistema e o fornecedor."},
        {"Etapa": 3, "Orientação": "Na aba do fornecedor, informe a quantidade em SOMA PENDENTE FORNECEDOR."},
        {"Etapa": 4, "Orientação": "Na aba do sistema, informe a quantidade em SOMA PENDENTE."},
        {"Etapa": 5, "Orientação": "Pode haver referências repetidas; o NEXO soma automaticamente antes de comparar."},
        {"Etapa": 6, "Orientação": "O NEXO identifica: quantidade divergente, só no fornecedor, só no sistema e OK."},
    ])

    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        supplier.to_excel(writer, sheet_name="PENDENCIA FORNECEDOR SOMA", index=False)
        system.to_excel(writer, sheet_name="PENDENCIA SISTEMA SOMA", index=False)
        instructions.to_excel(writer, sheet_name="INSTRUCOES", index=False)
        example_supplier.to_excel(writer, sheet_name="EXEMPLO FORNECEDOR", index=False)
        example_system.to_excel(writer, sheet_name="EXEMPLO SISTEMA", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(name="Times New Roman", bold=True, color=WHITE)
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Times New Roman", size=11)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                max_len = max([len(str(cell.value or "")) for cell in col[:200]] + [10])
                ws.column_dimensions[letter].width = min(max(max_len + 2, 14), 46)
    out.seek(0)
    return out.getvalue()


def comparison_xlsx_bytes(result, system_base, supplier_base):
    summary = (
        result.groupby("STATUS", as_index=False)
        .agg(
            REFERENCIAS=("REFERENCIA", "count"),
            QTD_SISTEMA=("QTD SISTEMA", "sum"),
            QTD_FORNECEDOR=("QTD FORNECEDOR", "sum"),
        )
    )
    summary["DIFERENÇA"] = summary["QTD_FORNECEDOR"] - summary["QTD_SISTEMA"]
    summary["_ORDEM"] = summary["STATUS"].map({s: i for i, s in enumerate(STATUS_ORDER)}).fillna(99)
    summary = summary.sort_values("_ORDEM").drop(columns="_ORDEM")

    divergences = result[result["STATUS"].ne("OK")].copy()

    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="RESUMO", index=False)
        divergences.to_excel(writer, sheet_name="DIVERGENCIAS", index=False)
        result.to_excel(writer, sheet_name="TODOS", index=False)
        system_base.to_excel(writer, sheet_name="BASE SISTEMA", index=False)
        supplier_base.to_excel(writer, sheet_name="BASE FORNECEDOR", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(name="Times New Roman", bold=True, color=WHITE)
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Times New Roman", size=11)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                max_len = max([len(str(cell.value or "")) for cell in col[:400]] + [10])
                ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 48)

            status_col = None
            for cell in ws[1]:
                if str(cell.value).strip().upper() == "STATUS":
                    status_col = cell.column
                    break
            if status_col:
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=status_col)
                    fill = STATUS_FILL.get(str(cell.value), None)
                    if fill:
                        cell.fill = PatternFill("solid", fgColor=fill)
                        cell.font = Font(name="Times New Roman", bold=True)

    out.seek(0)
    return out.getvalue()


def render_pendencias():
    st.title("🔎 Confronto de Pendências")
    st.caption(
        "Compare a carteira pendente do seu sistema com a carteira do fornecedor e identifique automaticamente todas as divergências."
    )

    with st.expander("📘 O que o NEXO vai conferir", expanded=False):
        st.markdown(
            """
            O confronto usa a **REFERÊNCIA** do produto e classifica cada item em quatro situações:

            - **QUANTIDADE DIVERGENTE** — aparece nos dois, mas com quantidades diferentes;
            - **SÓ NO FORNECEDOR** — existe na carteira do fornecedor, mas não no seu sistema;
            - **SÓ NO SISTEMA** — existe no seu sistema, mas não na carteira do fornecedor;
            - **OK** — aparece nos dois e a quantidade confere.

            Referências repetidas são somadas antes da comparação.
            """
        )

    mode = st.radio(
        "Como deseja importar?",
        ["Arquivo único com duas abas", "Dois arquivos separados"],
        horizontal=True,
        key="pending_import_mode",
    )

    col_upload, col_model = st.columns([2.2, 1])
    with col_model:
        st.markdown("**Primeira vez usando?**")
        st.download_button(
            "📥 Baixar modelo de pendências",
            data=pending_template_bytes(),
            file_name="NEXO_modelo_confronto_pendencias.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            key="pending_model_download",
        )

    system_df = None
    supplier_df = None

    try:
        with col_upload:
            if mode == "Arquivo único com duas abas":
                combined = st.file_uploader(
                    "Importe o arquivo com as pendências do SISTEMA e do FORNECEDOR",
                    type=["xlsx", "xls"],
                    key="pending_combined_file",
                    help="O NEXO procura automaticamente uma aba com 'FORNECEDOR' e outra com 'SISTEMA' no nome.",
                )
                if combined:
                    system_df, supplier_df, system_sheet, supplier_sheet = read_combined_workbook(combined)
                    st.success(
                        f"Abas reconhecidas: Sistema = {system_sheet} | Fornecedor = {supplier_sheet}"
                    )
            else:
                system_file = st.file_uploader(
                    "1) Pendências do seu SISTEMA",
                    type=["xlsx", "xls", "csv"],
                    key="pending_system_file",
                )
                supplier_file = st.file_uploader(
                    "2) Pendências do FORNECEDOR",
                    type=["xlsx", "xls", "csv"],
                    key="pending_supplier_file",
                )
                if system_file and supplier_file:
                    system_df = read_single_table(system_file)
                    supplier_df = read_single_table(supplier_file)
    except Exception as exc:
        st.error(str(exc))
        return

    if system_df is None or supplier_df is None:
        st.info("Importe as duas bases para iniciar o confronto.")
        return

    try:
        result = compare_pending(system_df, supplier_df)
        system_base = prepare_pending(system_df, "SISTEMA").rename(columns={"QUANTIDADE": "QTD SISTEMA"})
        supplier_base = prepare_pending(supplier_df, "FORNECEDOR").rename(columns={"QUANTIDADE": "QTD FORNECEDOR"})
    except Exception as exc:
        st.error(str(exc))
        return

    total = len(result)
    ok_count = int((result["STATUS"] == "OK").sum())
    qty_diff_count = int((result["STATUS"] == "QUANTIDADE DIVERGENTE").sum())
    supplier_only_count = int((result["STATUS"] == "SÓ NO FORNECEDOR").sum())
    system_only_count = int((result["STATUS"] == "SÓ NO SISTEMA").sum())
    divergent_count = total - ok_count

    st.markdown("### Resumo do confronto")
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Referências", f"{total:,}".replace(",", "."))
    c2.metric("Divergências", f"{divergent_count:,}".replace(",", "."))
    c3.metric("Qtd. diferente", f"{qty_diff_count:,}".replace(",", "."))
    c4.metric("Só fornecedor", f"{supplier_only_count:,}".replace(",", "."))
    c5.metric("Só sistema", f"{system_only_count:,}".replace(",", "."))

    q1, q2, q3 = st.columns(3)
    q1.metric("Quantidade total no sistema", f"{result['QTD SISTEMA'].sum():,.0f}".replace(",", "."))
    q2.metric("Quantidade total no fornecedor", f"{result['QTD FORNECEDOR'].sum():,.0f}".replace(",", "."))
    q3.metric("Itens OK", f"{ok_count:,}".replace(",", "."))

    tab_div, tab_all, tab_summary = st.tabs(["⚠️ Divergências", "📋 Todos os itens", "📊 Resumo"])

    with tab_div:
        statuses = ["QUANTIDADE DIVERGENTE", "SÓ NO FORNECEDOR", "SÓ NO SISTEMA"]
        f1, f2 = st.columns([2, 1])
        selected = f1.multiselect(
            "Filtrar situação",
            options=statuses,
            default=statuses,
            key="pending_status_filter",
        )
        search = f2.text_input("Buscar referência", key="pending_search_ref").strip().upper()

        view = result[result["STATUS"].isin(selected)].copy()
        if search:
            view = view[view["REFERENCIA"].str.contains(search, case=False, na=False)]

        st.dataframe(
            view[[
                "REFERENCIA",
                "QTD SISTEMA",
                "QTD FORNECEDOR",
                "DIFERENÇA FORNECEDOR - SISTEMA",
                "STATUS",
                "LEITURA",
            ]],
            use_container_width=True,
            hide_index=True,
        )
        st.caption(
            "Diferença = quantidade do fornecedor − quantidade do sistema. Valor positivo significa que o fornecedor tem mais; negativo significa que o sistema tem mais."
        )

    with tab_all:
        st.dataframe(
            result[[
                "REFERENCIA",
                "QTD SISTEMA",
                "QTD FORNECEDOR",
                "DIFERENÇA FORNECEDOR - SISTEMA",
                "STATUS",
                "LEITURA",
            ]],
            use_container_width=True,
            hide_index=True,
        )

    with tab_summary:
        summary = (
            result.groupby("STATUS", as_index=False)
            .agg(
                REFERENCIAS=("REFERENCIA", "count"),
                QTD_SISTEMA=("QTD SISTEMA", "sum"),
                QTD_FORNECEDOR=("QTD FORNECEDOR", "sum"),
            )
        )
        summary["DIFERENÇA"] = summary["QTD_FORNECEDOR"] - summary["QTD_SISTEMA"]
        summary["_ORDEM"] = summary["STATUS"].map({s: i for i, s in enumerate(STATUS_ORDER)}).fillna(99)
        summary = summary.sort_values("_ORDEM").drop(columns="_ORDEM")
        st.dataframe(summary, use_container_width=True, hide_index=True)

    st.markdown("### Exportar resultado")
    st.download_button(
        "📥 Baixar confronto completo em Excel",
        data=comparison_xlsx_bytes(result, system_base, supplier_base),
        file_name="NEXO_confronto_pendencias.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="pending_export_xlsx",
    )
    st.caption(
        "O Excel exportado contém RESUMO, DIVERGENCIAS, TODOS, BASE SISTEMA e BASE FORNECEDOR."
    )
